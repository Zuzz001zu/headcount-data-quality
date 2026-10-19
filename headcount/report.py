"""Excel report for the HR statistics team (in French, like the source data)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from headcount.pipeline import Result
from headcount.rules import RULES

MONTHS = [
    "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
]  # fmt: skip

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
SECTION_FONT = Font(bold=True, size=11, color="1F3864")
ERROR_FILL = PatternFill("solid", fgColor="F8D7DA")
WARNING_FILL = PatternFill("solid", fgColor="FFF3CD")
OK_FILL = PatternFill("solid", fgColor="D4EDDA")

CHF = "#,##0.00"
ONE_DECIMAL = "0.0"

HEADERS = {
    "department": "Département",
    "unit_code": "Code unité",
    "unit_name": "Unité",
    "headcount": "Effectif",
    "fte": "ETP",
    "women": "Femmes",
    "men": "Hommes",
    "women_pct": "Femmes (%)",
    "average_age": "Âge moyen",
    "payroll_total": "Masse salariale (CHF)",
    "average_activity_rate": "Taux d'activité moyen (%)",
    "contract_type": "Type de contrat",
    "share_pct": "Part (%)",
    "age_band": "Tranche d'âge",
    "rule_id": "Règle",
    "severity": "Gravité",
    "source": "Source",
    "employee_id": "Matricule",
    "field": "Champ",
    "value": "Valeur",
    "description": "Description",
    "entries": "Entrées",
    "exits": "Sorties",
    "net_change": "Variation nette",
    "exit_rate_pct": "Taux de sortie (%)",
    "movement": "Mouvement",
    "date": "Date",
    "last_name": "Nom",
    "first_name": "Prénom",
}
FORMATS = {
    "fte": ONE_DECIMAL,
    "women_pct": ONE_DECIMAL,
    "average_age": ONE_DECIMAL,
    "average_activity_rate": ONE_DECIMAL,
    "share_pct": ONE_DECIMAL,
    "exit_rate_pct": ONE_DECIMAL,
    "payroll_total": CHF,
}
SEVERITY_LABELS = {"error": "Bloquante", "warning": "Avertissement"}
SERIES_COLORS = ["1F3864", "8FAADC"]


def period_label(period: pd.Period) -> str:
    return f"{MONTHS[period.month - 1]} {period.year}"


def write_table(ws: Worksheet, df: pd.DataFrame, top: int, left: int = 1) -> int:
    """Write df with a styled header row. Returns the last row used."""
    for j, column in enumerate(df.columns):
        cell = ws.cell(row=top, column=left + j, value=HEADERS.get(column, column))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, record in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(record):
            if pd.isna(value):
                value = None
            elif hasattr(value, "item"):
                value = value.item()
            cell = ws.cell(row=top + i, column=left + j, value=value)
            if df.columns[j] in FORMATS:
                cell.number_format = FORMATS[df.columns[j]]
    return top + len(df)


def _is_free_text(cell) -> bool:
    """Titles, section headings and notes may overflow; they should not widen columns."""
    color = cell.font.color.rgb if cell.font.color is not None else None
    return bool(cell.font.i) or (isinstance(color, str) and color.endswith("1F3864"))


def print_setup(ws: Worksheet) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def autosize(ws: Worksheet, minimum: int = 10, maximum: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or _is_free_text(cell):
                continue
            length = len(str(cell.value))
            if isinstance(cell.value, float):
                length = len(f"{cell.value:,.2f}")
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = max(
            minimum, min(maximum, width + 3)
        )


def summary_sheet(ws: Worksheet, result: Result) -> None:
    s = result.sources
    ws["A1"] = "Effectifs et qualité des données"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Période : {period_label(s.period)} — effectif au {s.reference_date:%d.%m.%Y}"

    t = result.totals
    moves = result.tables["movements_by_department"]
    figures = [
        ("Lignes SIRH lues", len(s.employees), None),
        ("Lignes de salaire lues", len(s.payroll), None),
        ("Anomalies bloquantes", result.errors, None),
        ("Avertissements", result.warnings, None),
        ("Collaborateurs exclus des statistiques", result.excluded_employees, None),
        ("Effectif retenu", int(t["headcount"]), None),
        ("Équivalents plein temps (ETP)", float(t["fte"]), ONE_DECIMAL),
        ("Taux d'activité moyen (%)", float(t["average_activity_rate"]), ONE_DECIMAL),
        ("Masse salariale mensuelle (CHF)", float(t["payroll_total"]), CHF),
        ("Entrées sur la période", int(moves["entries"].sum()), None),
        ("Sorties sur la période", int(moves["exits"].sum()), None),
    ]
    ws["A4"] = "Chiffres clés"
    ws["A4"].font = SECTION_FONT
    for i, (label, value, fmt) in enumerate(figures, start=5):
        ws.cell(row=i, column=1, value=label)
        cell = ws.cell(row=i, column=2, value=value)
        cell.font = Font(bold=True)
        if fmt:
            cell.number_format = fmt

    top = 5 + len(figures) + 1
    ws.cell(row=top, column=1, value="Contrôles appliqués").font = SECTION_FONT
    counts = result.issues.groupby("rule_id").size()
    checks = pd.DataFrame(
        {
            "description": [r.description for r in RULES],
            "rule_id": [r.id for r in RULES],
            "severity": [SEVERITY_LABELS[r.severity] for r in RULES],
            "source": [r.source for r in RULES],
            "Occurrences": [int(counts.get(r.id, 0)) for r in RULES],
        }
    )
    last = write_table(ws, checks, top + 1)
    for row in range(top + 2, last + 1):
        cell = ws.cell(row=row, column=5)
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="center")
        if cell.value == 0:
            cell.fill = OK_FILL
        else:
            cell.fill = (
                ERROR_FILL if ws.cell(row=row, column=3).value == "Bloquante" else WARNING_FILL
            )

    ws.cell(
        row=last + 2,
        column=1,
        value=(
            "Les collaborateurs concernés par une anomalie bloquante sont exclus des statistiques "
            "jusqu'à correction dans le système source. Détail dans l'onglet « Anomalies »."
        ),
    ).font = Font(italic=True, color="595959")


def bar_chart(title: str, y_title: str, height: float = 8, width: float = 18) -> BarChart:
    chart = BarChart()
    chart.title = title
    chart.y_axis.title = y_title
    chart.height = height
    chart.width = width
    chart.y_axis.majorGridlines = None
    return chart


def color_series(chart: BarChart) -> None:
    for series, color in zip(chart.series, SERIES_COLORS, strict=False):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color


def headcount_sheet(ws: Worksheet, result: Result) -> None:
    ws["A1"] = "Effectifs par département"
    ws["A1"].font = TITLE_FONT
    departments = result.tables["headcount_by_department"]
    last = write_table(ws, departments, 3)

    chart = bar_chart("Effectif et ETP par département", "Personnes")
    data = Reference(ws, min_col=2, max_col=3, min_row=3, max_row=last)
    labels = Reference(ws, min_col=1, min_row=4, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    color_series(chart)
    ws.add_chart(chart, f"A{last + 3}")


def units_sheet(ws: Worksheet, result: Result) -> None:
    ws["A1"] = "Effectifs par unité"
    ws["A1"].font = TITLE_FONT
    write_table(ws, result.tables["headcount_by_unit"], 3)
    ws.freeze_panes = "A4"


def breakdown_sheet(ws: Worksheet, result: Result) -> None:
    ws["A1"] = "Répartition de l'effectif"
    ws["A1"].font = TITLE_FONT
    ws["A3"] = "Par type de contrat"
    ws["A3"].font = SECTION_FONT
    last = write_table(ws, result.tables["headcount_by_contract"], 4)

    top = last + 2
    ws.cell(row=top, column=1, value="Par tranche d'âge et sexe").font = SECTION_FONT
    ages = result.tables["headcount_by_age"]
    last = write_table(ws, ages, top + 1)

    chart = bar_chart("Effectif par tranche d'âge et sexe", "Personnes", width=14)
    data = Reference(ws, min_col=2, max_col=3, min_row=top + 1, max_row=last)
    labels = Reference(ws, min_col=1, min_row=top + 2, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labels)
    color_series(chart)
    ws.add_chart(chart, "F3")


def movements_sheet(ws: Worksheet, result: Result) -> None:
    ws["A1"] = "Entrées et sorties"
    ws["A1"].font = TITLE_FONT
    last = write_table(ws, result.tables["movements_by_department"], 3)

    top = last + 2
    ws.cell(row=top, column=1, value="Détail des mouvements").font = SECTION_FONT
    people = result.tables["movements_list"].copy()
    people["date"] = pd.to_datetime(people["date"]).dt.strftime("%d.%m.%Y")
    write_table(ws, people, top + 1)


def issues_sheet(ws: Worksheet, result: Result) -> None:
    # "error" sorts before "warning", so blocking anomalies come first
    issues = result.issues.sort_values(["severity", "rule_id", "employee_id"], na_position="last")
    issues = issues.assign(severity=issues["severity"].map(SEVERITY_LABELS))
    last = write_table(ws, issues, 1)
    for row in range(2, last + 1):
        fill = ERROR_FILL if ws.cell(row=row, column=2).value == "Bloquante" else WARNING_FILL
        ws.cell(row=row, column=2).fill = fill
    ws.auto_filter.ref = f"A1:{get_column_letter(len(issues.columns))}{max(last, 1)}"
    ws.freeze_panes = "A2"


def write_report(result: Result, out_dir: Path) -> Path:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    period = result.sources.period

    wb = Workbook()
    sheets = [
        ("Synthèse", summary_sheet),
        ("Effectifs", headcount_sheet),
        ("Unités", units_sheet),
        ("Répartition", breakdown_sheet),
        ("Mouvements", movements_sheet),
        ("Anomalies", issues_sheet),
    ]
    for index, (title, build) in enumerate(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        build(ws, result)
        autosize(ws)
        print_setup(ws)

    path = out_dir / f"rapport_effectifs_{period}.xlsx"
    wb.save(path)
    return path


def write_issues_csv(result: Result, out_dir: Path) -> Path:
    """Flat list of anomalies, to send back to the owners of the source systems."""
    path = Path(out_dir) / f"anomalies_{result.sources.period}.csv"
    result.issues.to_csv(path, sep=";", index=False, encoding="utf-8-sig")
    return path
