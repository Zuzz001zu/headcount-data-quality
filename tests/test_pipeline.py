"""End-to-end: generated exports -> checks -> Excel report."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from headcount import cli, pipeline, report
from headcount.generate import generate, write_sources
from tests.conftest import PERIOD


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    raw = tmp_path_factory.mktemp("raw")
    data = generate(PERIOD, size=400, seed=7)
    write_sources(data, raw, PERIOD)
    return raw, data


def test_every_injected_anomaly_is_found(generated):
    raw, data = generated
    result = pipeline.run(raw, PERIOD)
    found = result.issues.groupby("rule_id").size().to_dict()
    assert found == data.injected


def test_cosmetic_differences_are_not_reported_as_errors(generated):
    raw, data = generated
    result = pipeline.run(raw, PERIOD)
    assert result.errors == sum(
        n for rule, n in data.injected.items() if rule not in {"HR-09", "PAY-05", "PAY-06"}
    )


def test_generation_is_reproducible():
    first = generate(PERIOD, size=200, seed=1)
    second = generate(PERIOD, size=200, seed=1)
    pd.testing.assert_frame_equal(first.employees, second.employees)
    pd.testing.assert_frame_equal(first.payroll, second.payroll)


def test_generator_refuses_too_small_samples():
    with pytest.raises(ValueError, match="at least"):
        generate(PERIOD, size=50)


def test_report_has_all_sheets(generated, tmp_path):
    raw, _ = generated
    result = pipeline.run(raw, PERIOD)
    path = report.write_report(result, tmp_path)

    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Synthèse",
        "Effectifs",
        "Unités",
        "Répartition",
        "Mouvements",
        "Anomalies",
    ]
    assert wb["Synthèse"]["A2"].value.startswith("Période : septembre 2026")
    assert wb["Anomalies"].max_row == len(result.issues) + 1
    assert len(wb["Effectifs"]._charts) == 1


def test_cli_run(generated, tmp_path, capsys):
    raw, _ = generated
    code = cli.main(["run", "--input", str(raw), "--period", "2026-09", "--out", str(tmp_path)])

    assert code == 0
    assert (tmp_path / "rapport_effectifs_2026-09.xlsx").exists()
    assert (tmp_path / "anomalies_2026-09.csv").exists()
    assert "Headcount on 2026-09-30" in capsys.readouterr().out


def test_cli_can_fail_on_blocking_anomalies(generated, tmp_path):
    raw, _ = generated
    args = ["run", "--input", str(raw), "--period", "2026-09", "--out", str(tmp_path)]
    assert cli.main([*args, "--fail-on-errors"]) == 1


def test_cli_reports_missing_files(tmp_path, capsys):
    code = cli.main(["run", "--input", str(tmp_path), "--period", "2026-09"])
    assert code == 2
    assert "not found" in capsys.readouterr().err


def test_cli_rejects_bad_period():
    with pytest.raises(SystemExit):
        cli.main(["run", "--period", "septembre"])


def test_blocking_anomalies_are_listed_first(generated, tmp_path):
    raw, _ = generated
    path = report.write_report(pipeline.run(raw, PERIOD), tmp_path)
    ws = load_workbook(path)["Anomalies"]
    severities = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert severities == sorted(severities, key=lambda s: s != "Bloquante")
